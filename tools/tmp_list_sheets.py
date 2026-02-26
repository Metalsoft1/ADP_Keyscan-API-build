from openpyxl import load_workbook

wb = load_workbook("samples/dataset.xlsx", read_only=True, data_only=True)
print("Sheets:")
for s in wb.sheetnames:
    print(" -", s)
