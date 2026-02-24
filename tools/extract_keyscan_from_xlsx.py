from openpyxl import load_workbook
import csv
from pathlib import Path

xlsx_path = Path("samples/dataset.xlsx")
out_csv = Path("samples/keyscan_export.csv")

wb = load_workbook(xlsx_path, read_only=True, data_only=True)
ws = wb.active  # Foglio1

# Read all rows, keep as lists, trim trailing empty cells
rows = []
for r in ws.iter_rows(values_only=True):
    r = list(r)
    # trim trailing Nones
    while r and (r[-1] is None or str(r[-1]).strip() == ""):
        r.pop()
    rows.append(r)

# Remove fully empty rows
rows = [r for r in rows if any((c is not None and str(c).strip() != "") for c in r)]

# Find the first row that looks like a header for data:
# must contain at least one of these key columns (case-insensitive)
needles = ["credential", "credential_number", "datetime", "date", "location", "door", "employee", "badge", "card"]
header_idx = None
for i, r in enumerate(rows):
    joined = " | ".join(str(c).strip().lower() for c in r if c is not None)
    if any(n in joined for n in needles) and "select * from" not in joined and "query" not in joined:
        header_idx = i
        break

if header_idx is None:
    print("ERROR: Could not find a data header row in the XLSX.")
    print("Top non-empty rows found (first 10):")
    for r in rows[:10]:
        print(r)
    raise SystemExit(2)

header = [str(c).strip() if c is not None else "" for c in rows[header_idx]]
data = rows[header_idx+1:]

# Normalize row lengths to header length
h_len = len(header)
norm = []
for r in data:
    r = list(r)
    if len(r) < h_len:
        r += [""] * (h_len - len(r))
    elif len(r) > h_len:
        r = r[:h_len]
    norm.append([("" if c is None else str(c)) for c in r])

out_csv.parent.mkdir(parents=True, exist_ok=True)
with out_csv.open("w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow(header)
    w.writerows(norm)

print(f"Wrote: {out_csv} (rows={len(norm)}, cols={len(header)})")
print("Header:")
print(header)
